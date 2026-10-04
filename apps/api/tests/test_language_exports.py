"""A deck exported in another language, read back by readers that are not us.

Integration plan 01 §3.10: the export takes a language, applies that language's
overlay before anything is measured, embeds the faces that draw its script, and
— for PowerPoint — carries the recordings in that language. `python-pptx` and
`pypdf` are independent implementations of the formats, so a file they read is
a file PowerPoint and a PDF viewer read.

Runs the real exporter CLI (`apps/worker/src/cli.ts`), the path an export job
takes, which needs Node and a Chromium for text measurement.
"""

from __future__ import annotations

import base64
import json
import math
import re
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path

import pytest

from deckastra_api import audio

ROOT = Path(__file__).resolve().parents[3]
FIXTURE = ROOT / "packages/presentation-schema/fixtures/multilingual-narrated.mydeck.json"
CLI = ROOT / "apps/worker/src/cli.ts"

pptx = pytest.importorskip("pptx", reason="python-pptx is the independent reader this test needs")
pypdf = pytest.importorskip("pypdf", reason="pypdf is the independent reader this test needs")

if shutil.which("npx") is None:  # pragma: no cover - environment guard
    pytest.skip("npx is needed to run the exporter", allow_module_level=True)

pytestmark = pytest.mark.slow


def _export(
    tmp_path: Path, kind: str, locale: str | None, with_audio: bool, document: dict | None = None
) -> tuple[Path, dict]:
    document = document or json.loads(FIXTURE.read_text(encoding="utf-8"))
    assets = []
    if with_audio:
        for asset in document["assets"]:
            if asset["type"] != "audio":
                continue
            seconds = asset.get("durationMs", 1000) / 1000
            samples = [0.2 * math.sin(2 * math.pi * 330 * i / 8000) for i in range(int(seconds * 8000))]
            assets.append(
                {
                    "assetId": asset["id"],
                    "storageKey": asset["storageKey"],
                    "mimeType": "audio/wav",
                    "data": base64.b64encode(audio.encode_wav(samples, 8000)).decode("ascii"),
                }
            )
    document_path = tmp_path / "deck.json"
    document_path.write_text(json.dumps(document), encoding="utf-8")
    assets_path = tmp_path / "assets.json"
    assets_path.write_text(json.dumps(assets), encoding="utf-8")
    output = tmp_path / f"deck.{kind}"
    invocation = {
        "kind": kind,
        "output": str(output),
        "documentPath": str(document_path),
        "assetsPath": str(assets_path),
        "options": {"locale": locale} if locale else {},
    }
    finished = subprocess.run(
        ["npx", "tsx", str(CLI)],
        input=json.dumps(invocation),
        cwd=ROOT,
        capture_output=True,
        text=True,
        encoding="utf-8",
        timeout=300,
        shell=sys.platform == "win32",
    )
    answer = json.loads(finished.stdout.strip().splitlines()[-1]) if finished.stdout.strip() else {}
    if not answer.get("ok"):
        pytest.fail(f"the exporter failed: {answer or finished.stderr[-2000:]}")
    return output, answer


def test_a_hindi_powerpoint_has_hindi_words_and_hindi_narration(tmp_path):
    from pptx import Presentation

    path, answer = _export(tmp_path, "pptx", "hi-IN", with_audio=True)
    assert answer["filename"] == "Multilingual-Narrated-Deck-hi-IN.pptx"
    deck = Presentation(str(path))
    texts = [shape.text_frame.text for shape in deck.slides[0].shapes if shape.has_text_frame]
    assert any("एक डेक, हर भाषा" in text for text in texts), texts
    assert not any("One deck, every language" in text for text in texts)

    # The recordings are parts in the package, the relationship graph python-pptx
    # resolved on open reaches them, and only the Hindi ones are there.
    with zipfile.ZipFile(path) as package:
        media = [name for name in package.namelist() if name.startswith("ppt/media/media")]
        slide = package.read("ppt/slides/slide2.xml").decode("utf-8")
    assert len(media) == 5  # four Hindi lines and the pop
    assert slide.count("<a:audioFile ") == 5
    assert slide.count('cmd="playFrom(0.0)"') == 4 + 3


def test_a_hindi_pdf_embeds_the_face_that_draws_devanagari(tmp_path):
    from pypdf import PdfReader

    path, answer = _export(tmp_path, "pdf", "hi-IN", with_audio=False)
    reader = PdfReader(str(path))
    assert len(reader.pages) == 3
    fonts = _pdf_fonts(reader)
    assert any("NotoSansDevanagari" in name.replace("-", "").replace(" ", "") for name in fonts), fonts
    # Static faces, embedded as TrueType (Type0). A variable face became Type3,
    # whose character map is the cmap alone, and its shaped glyphs read back
    # as U+0000 (plan 01 recheck, 2026-10-03).
    subtypes = {
        str(font.get_object().get("/Subtype"))
        for page in reader.pages
        for font in ((page.get("/Resources") or {}).get("/Font") or {}).values()
    }
    assert subtypes == {"/Type0"}, subtypes
    # Not the machine's own Devanagari face: an export host has only what ships.
    assert not any("Nirmala" in name or "Mangal" in name for name in fonts), fonts
    # Every glyph has its characters, for a reader that reads only the map
    # (pypdf), including the vowel sign shaping turned into a width variant.
    text = "".join(page.extract_text() for page in reader.pages)
    assert "\x00" not in text, repr(text)
    normalized = re.sub(r"\s+([,।])", r"\1", " ".join(text.split()))
    assert "एक डेक, हर भाषा" in normalized
    assert "अनूदित शब्द" in normalized  # exact logical cluster order, not अनूिदत
    # And the report says the narration could not come with it.
    assert any(w["feature"] == "audio" and w["action"] == "dropped" for w in answer["report"]["warnings"])


def test_the_original_language_needs_no_overlay(tmp_path):
    from pptx import Presentation

    path, answer = _export(tmp_path, "pptx", None, with_audio=True)
    assert answer["filename"] == "Multilingual-Narrated-Deck.pptx"
    texts = [shape.text_frame.text for shape in Presentation(str(path)).slides[0].shapes if shape.has_text_frame]
    assert any("One deck, every language" in text for text in texts)


def _pdf_fonts(reader) -> set[str]:
    fonts: set[str] = set()
    for page in reader.pages:
        resources = page.get("/Resources") or {}
        for font in (resources.get("/Font") or {}).values():
            font = font.get_object()
            descriptor = font.get("/FontDescriptor")
            name = descriptor.get_object().get("/FontName") if descriptor else font.get("/BaseFont")
            fonts.add(str(name))
            for child in font.get("/DescendantFonts") or []:
                child_descriptor = child.get_object().get("/FontDescriptor")
                if child_descriptor:
                    fonts.add(str(child_descriptor.get_object().get("/FontName")))
    return fonts


@pytest.mark.skipif(
    sys.platform != "win32" or not Path("C:/Windows/Fonts/YuGothM.ttc").exists(),
    reason="checks the Windows CJK fallback; no Noto CJK face ships yet (plan 01 §3.9 defers them to packs)",
)
def test_a_japanese_pdf_draws_real_glyphs_from_the_systems_cjk_face(tmp_path):
    """CJK faces are not bundled (they are the plan's downloadable packs), so a
    Japanese deck draws in the operating system's face. What must hold is that it
    *draws*: real glyphs from a CJK face in the PDF, extractable as Japanese, not
    boxes from a Latin one."""
    from pypdf import PdfReader

    document = json.loads(FIXTURE.read_text(encoding="utf-8"))
    document.pop("locales", None)
    document["metadata"]["language"] = "ja"
    title = document["slides"][0]["elements"][0]
    title["content"]["blocks"] = [{**title["content"]["blocks"][0], "spans": [{"text": "ひとつのデッキ、すべての言語"}]}]

    path, _answer = _export(tmp_path, "pdf", None, with_audio=False, document=document)
    reader = PdfReader(str(path))
    assert "ひとつのデッキ" in reader.pages[0].extract_text()
    fonts = _pdf_fonts(reader)
    assert any(name for name in fonts if any(face in name.replace("-", "").replace(" ", "") for face in ("YuGothic", "Meiryo", "NotoSansJP", "MSGothic"))), fonts


@pytest.mark.skipif(shutil.which("pdftotext") is None, reason="Poppler's pdftotext reads the /ActualText Chromium writes")
def test_an_arabic_pdf_reads_back_as_its_own_letters(tmp_path):
    """Read complete source paragraphs through Poppler's ActualText support.
    Shaped dotted letters must retain their identity and no null glyphs may
    leak into extraction. Poppler wraps visual-order RTL runs in RLE/PDF;
    decode those explicit runs before comparing logical Unicode words."""
    path, _answer = _export(tmp_path, "pdf", "ar", with_audio=False)
    from pypdf import PdfReader

    # Include click-reveal text, which Chromium may paint in Form XObjects.
    assert all("\x00" not in page.extract_text() for page in PdfReader(path).pages)
    finished = subprocess.run(
        ["pdftotext", "-l", "1", "-enc", "UTF-8", str(path), "-"], capture_output=True, check=True
    )
    text = finished.stdout.decode("utf-8")
    letters = "".join(char for char in text if "\u0600" <= char <= "\u06ff")
    expected = "".join(char for char in "عرض واحد، كل لغة" if "\u0600" <= char <= "\u06ff")
    # Order within the line is the reader's bidi reconstruction; the letters are what the file carries.
    assert sorted(letters) == sorted(expected) or expected in letters, (text, letters)
    assert "ض" in letters
    # Complete logical words/paragraphs, not only the right bag of glyphs.
    logical = re.sub("\u202b(.*?)\u202c", lambda match: match[1][::-1], text, flags=re.DOTALL)
    normalized = " ".join(logical.translate(dict.fromkeys(map(ord, "\u202a\u202b\u202c\u202d\u202e\u2066\u2067\u2068\u2069"))).split())
    assert "عرض واحد، كل لغة" in normalized, text
    assert "\x00" not in text
