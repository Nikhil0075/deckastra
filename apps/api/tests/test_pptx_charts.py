"""A bubble map reopens in a PowerPoint reader that is not ours (2026-10-13).

The land is one custom outline with a move for every landmass, the largest
custom shape this exporter writes. `python-pptx` must open the file and find,
inside the chart's group, that outline and an oval for every bubble.
"""

from __future__ import annotations

import shutil
import subprocess
import sys

import pytest

from tests.test_pptx_opens import ROOT, FIXTURE

pptx = pytest.importorskip("pptx", reason="python-pptx is the independent reader this test needs")

if shutil.which("npx") is None:  # pragma: no cover - environment guard
    pytest.skip("npx is needed to run the exporter", allow_module_level=True)

EMIT = """
import { writeFileSync, readFileSync } from "node:fs";
import { buildDocumentScene } from "@deckastra/renderer";
import { fontManifest } from "@deckastra/export-core";
import { buildPptx } from "@deckastra/export-pptx";

const doc = JSON.parse(readFileSync(process.argv[2], "utf8"));
doc.slides[0].elements.push({
  id: "el_01JB8Z9K2QW4RN7F3XG5HTMDZZ", type: "chart", name: "Map", chartType: "bubbleMap",
  transform: { x: 100, y: 100, width: 1200, height: 600, rotation: 0 },
  data: { type: "inline", rows: [
    { city: "London", lat: 51.5, lon: -0.13, v: 42 },
    { city: "Tokyo", lat: 35.68, lon: 139.69, v: 47 },
    { city: "Lagos", lat: 6.52, lon: 3.38, v: 18 },
  ] },
  encoding: { category: "city", value: "v", latitude: "lat", longitude: "lon" },
  chartStyle: { mapExtent: "world" }, altText: "Revenue by city",
});
const scene = buildDocumentScene(doc);
const scenes = new Map(scene.slides.map((s) => [s.slideId, s]));
const { bytes } = buildPptx({ document: doc, scenes, fontManifest: fontManifest(scenes.values()), options: {} });
writeFileSync(process.argv[3], bytes);
"""


@pytest.mark.slow
def test_a_bubble_map_opens_with_its_land_and_bubbles(tmp_path):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE

    script = ROOT / "packages/export-pptx/.emit-bubble-map.mts"
    script.write_text(EMIT, encoding="utf-8")
    out = tmp_path / "map.pptx"
    try:
        finished = subprocess.run(
            ["npx", "tsx", str(script), str(FIXTURE), str(out)],
            cwd=ROOT, capture_output=True, text=True, timeout=180, shell=sys.platform == "win32",
        )
    finally:
        script.unlink(missing_ok=True)
    if finished.returncode != 0:
        pytest.fail(f"the exporter failed:\n{finished.stderr[-2000:]}")

    slide = Presentation(str(out)).slides[0]
    groups = [shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.GROUP]
    shapes = [inner for group in groups for inner in group.shapes]
    ovals = [shape for shape in shapes if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and shape.auto_shape_type == MSO_SHAPE.OVAL]
    assert len(ovals) == 3, "every bubble is an oval in the chart's group"
    freeforms = [shape for shape in shapes if shape.shape_type == MSO_SHAPE_TYPE.FREEFORM]
    moves = max(len(shape._element.findall(".//{http://schemas.openxmlformats.org/drawingml/2006/main}moveTo")) for shape in freeforms)
    assert moves > 20, "the land is one outline with a move per landmass"
