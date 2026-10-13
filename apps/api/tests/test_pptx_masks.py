"""A masked picture reopens in a PowerPoint reader that is not ours (2026-10-13).

`python-pptx` resolves the picture's relationship itself and reads its crop
shape from the shape tree. An ellipse must come back as PowerPoint's own oval
crop, which a person can still change there; an arch must come back as a
picture with custom geometry carrying real curves. A reader that refused the
file, or found a rectangle, would mean PowerPoint does too.
"""

from __future__ import annotations

import json
import shutil
import subprocess
import sys
from pathlib import Path

import pytest

from tests.test_pptx_opens import ROOT, FIXTURE, _png

pptx = pytest.importorskip("pptx", reason="python-pptx is the independent reader this test needs")

if shutil.which("npx") is None:  # pragma: no cover - environment guard
    pytest.skip("npx is needed to run the exporter", allow_module_level=True)

EMIT = """
import { writeFileSync, readFileSync } from "node:fs";
import { buildDocumentScene } from "@deckastra/renderer";
import { fontManifest } from "@deckastra/export-core";
import { buildPptx } from "@deckastra/export-pptx";

const doc = JSON.parse(readFileSync(process.argv[2], "utf8"));
const image = doc.slides.flatMap((slide) => slide.elements).find((element) => element.type === "image");
image.fit = "cover";
image.mask = JSON.parse(process.argv[5]);
const scene = buildDocumentScene(doc);
const scenes = new Map(scene.slides.map((s) => [s.slideId, s]));
const images = new Map([[image.assetId, { bytes: new Uint8Array(readFileSync(process.argv[4])), contentType: "image/png" }]]);
const { bytes, result } = buildPptx({ document: doc, scenes, fontManifest: fontManifest(scenes.values()), options: {}, images });
writeFileSync(process.argv[3], bytes);
process.stdout.write(JSON.stringify(result.report));
"""


def _export(directory: Path, mask: dict) -> tuple[Path, dict]:
    script = ROOT / "packages/export-pptx/.emit-masks.mts"
    script.write_text(EMIT, encoding="utf-8")
    out = directory / "masked.pptx"
    picture = _png(directory / "picture.png")
    try:
        finished = subprocess.run(
            ["npx", "tsx", str(script), str(FIXTURE), str(out), str(picture), json.dumps(mask)],
            cwd=ROOT,
            capture_output=True,
            text=True,
            timeout=180,
            shell=sys.platform == "win32",
        )
    finally:
        script.unlink(missing_ok=True)
    if finished.returncode != 0:
        pytest.fail(f"the exporter failed:\n{finished.stderr[-2000:]}")
    return out, json.loads(finished.stdout.strip().splitlines()[-1])


def _picture(path: Path):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    deck = Presentation(str(path))
    pictures = [shape for slide in deck.slides for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
    assert pictures, "the masked picture is not in the file"
    return pictures[0]


@pytest.mark.slow
def test_an_ellipse_mask_is_powerpoints_own_oval_crop(tmp_path):
    from pptx.enum.shapes import MSO_SHAPE

    path, report = _export(tmp_path, {"type": "ellipse"})
    picture = _picture(path)
    assert picture.auto_shape_type == MSO_SHAPE.OVAL
    assert picture.image.size == (8, 8)
    assert [w for w in report["warnings"] if w["feature"] == "image mask"] == []


@pytest.mark.slow
def test_an_arch_mask_is_custom_geometry_with_curves(tmp_path):
    path, report = _export(tmp_path, {"type": "shape", "shape": "arch"})
    picture = _picture(path)
    sp_pr = picture._element.spPr
    namespace = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    assert sp_pr.find(f"{namespace}custGeom") is not None
    assert sp_pr.find(f"{namespace}prstGeom") is None
    assert len(sp_pr.findall(f".//{namespace}cubicBezTo")) >= 2
    assert [w for w in report["warnings"] if w["feature"] == "image mask"] == []
