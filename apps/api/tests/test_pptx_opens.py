"""The exported .pptx opens in a real PowerPoint reader.

Every other assertion about the package is made by the code that wrote it, which
means they all share the same misunderstanding if there is one. This one is made
by `python-pptx`, an independent implementation of the same specification: it
walks the relationship graph, resolves the master and layout, and reads the shape
tree. A package it refuses is a package PowerPoint refuses.

It is in the Python suite rather than the TypeScript one because that is where
the independent reader lives. The exporter is invoked through `tsx`, so this also
covers the thing no unit test can: that the package is written correctly to a
file rather than only correct in memory.
"""

from __future__ import annotations

import json
import shutil
import subprocess
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[3]
FIXTURE = ROOT / "packages/presentation-schema/fixtures/technical-deck.mydeck.json"

pptx = pytest.importorskip("pptx", reason="python-pptx is the independent reader this test needs")

if shutil.which("npx") is None:  # pragma: no cover - environment guard
    pytest.skip("npx is needed to run the exporter", allow_module_level=True)


EMIT = """
import { writeFileSync, readFileSync } from "node:fs";
import { buildDocumentScene } from "@deckastra/renderer";
import { fontManifest } from "@deckastra/export-core";
import { buildPptx } from "@deckastra/export-pptx";

const doc = JSON.parse(readFileSync(process.argv[2], "utf8"));
// A note the reader can look for. The fixture carries none, and a test that
// only checked their absence passed against an exporter writing them into the
// wrong element entirely.
doc.slides[0].speakerNotes = ["Remember to mention the budget.", "And thank the team."].join(String.fromCharCode(10));
const scene = buildDocumentScene(doc);
const scenes = new Map(scene.slides.map((s) => [s.slideId, s]));

// The picture the deck cites, supplied the way the worker supplies it. The
// fixture's storage key names no file on disk — a manifest records what a deck
// refers to, not bytes — so the caller brings them, which is exactly the
// contract `ExportInput.images` describes.
const images = new Map();
if (process.argv[4]) {
  const assetId = doc.slides
    .flatMap((slide) => slide.elements)
    .find((element) => element.type === "image").assetId;
  images.set(assetId, {
    bytes: new Uint8Array(readFileSync(process.argv[4])),
    contentType: "image/png",
  });
}

const { bytes, result } = buildPptx({
  document: doc,
  scenes,
  fontManifest: fontManifest(scenes.values()),
  options: { includeNotes: true },
  images,
});
writeFileSync(process.argv[3], bytes);
process.stdout.write(JSON.stringify(result.report));
"""


def _png(path: Path) -> Path:
    """A real 8x8 PNG, written here rather than committed.

    A committed binary beside its only consumer is a fixture nobody can read a
    diff of; eight pixels of one colour say what they are in the code that makes
    them.
    """
    import struct
    import zlib

    # One filter byte (0 = none) then eight magenta pixels, per row.
    row = bytes([0]) + bytes([255, 0, 255]) * 8

    def chunk(kind: bytes, body: bytes) -> bytes:
        return (
            struct.pack(">I", len(body))
            + kind
            + body
            + struct.pack(">I", zlib.crc32(kind + body) & 0xFFFFFFFF)
        )

    path.write_bytes(
        bytes([137, 80, 78, 71, 13, 10, 26, 10])
        + chunk(b"IHDR", struct.pack(">IIBBBBB", 8, 8, 8, 2, 0, 0, 0))
        + chunk(b"IDAT", zlib.compress(row * 8))
        + chunk(b"IEND", b"")
    )
    return path


@pytest.fixture(scope="module")
def exported(tmp_path_factory) -> tuple[Path, dict]:
    """Run the real exporter and hand back the file plus its report."""
    # Inside the workspace, because the package imports resolve through it.
    script = ROOT / "packages/export-pptx/.emit.mts"
    script.write_text(EMIT, encoding="utf-8")
    directory = tmp_path_factory.mktemp("pptx")
    out = directory / "deck.pptx"
    picture = _png(directory / "picture.png")

    try:
        finished = subprocess.run(
            ["npx", "tsx", str(script), str(FIXTURE), str(out), str(picture)],
            cwd=ROOT,
            capture_output=True,
            text=True,
            timeout=180,
            shell=sys.platform == "win32",
        )
    finally:
        script.unlink(missing_ok=True)

    if finished.returncode != 0:
        pytest.fail(f"the exporter failed:\n{finished.stderr[-2000:]}")

    return out, json.loads(finished.stdout.strip().splitlines()[-1])


def test_the_picture_is_in_the_package_and_the_reader_finds_it(exported):
    """Doc 04 §33.4: a client opening this must see the photograph, not a box.

    Every other assertion about the picture is made by the code that embedded it.
    `python-pptx` resolves the relationship itself, so a picture it can find is a
    picture PowerPoint can find — and the bytes it hands back are read from the
    part, not from what the writer believed it wrote.
    """
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    path, report = exported
    deck = Presentation(str(path))

    pictures = [
        shape
        for slide in deck.slides
        for shape in slide.shapes
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
    ]
    assert pictures, "the deck exported no picture at all"

    picture = pictures[0]
    # The bytes came back out of the package intact, through a reader that is not
    # us. `image.blob` is read from the media part the relationship points at.
    assert picture.image.blob.startswith(bytes([137, 80, 78, 71]))
    assert picture.image.ext == "png"
    assert picture.image.size == (8, 8)

    # Positioned, not dumped at the origin.
    assert picture.left > 0 and picture.top > 0
    assert picture.left + picture.width <= deck.slide_width

    # And the report no longer claims it was dropped, because it was not.
    dropped = [
        warning
        for warning in report["warnings"]
        if warning["feature"] == "image" and warning["action"] == "dropped"
    ]
    assert dropped == []


def test_a_real_powerpoint_reader_opens_it(exported):
    from pptx import Presentation

    path, report = exported
    deck = Presentation(str(path))

    # python-pptx resolves the relationship graph on open, so getting this far
    # means every part a reader needs is present and reachable.
    assert len(deck.slides) == report["slideCount"]
    # 13.333in x 7.5in in EMU (doc 04 §33.1).
    assert (deck.slide_width, deck.slide_height) == (12_192_000, 6_858_000)


def test_the_text_arrives_editable_and_in_the_right_place(exported):
    """Doc 04 §33.4: editability is the point, not pixel parity."""
    from pptx import Presentation

    path, _ = exported
    deck = Presentation(str(path))
    first = deck.slides[0]

    texts = [shape.text_frame.text for shape in first.shapes if shape.has_text_frame]
    assert any("Agents propose" in text for text in texts)

    # Positioned in EMU derived from the document's own coordinates, so a
    # recipient sees the layout the author built rather than a reflow.
    boxes = [shape for shape in first.shapes if shape.has_text_frame]
    assert all(shape.left is not None and shape.top is not None for shape in boxes)
    assert all(0 <= shape.left < deck.slide_width for shape in boxes)


def test_shape_names_are_stable_and_derived_from_element_ids(exported):
    """Doc 04 §33.3: PowerPoint's Morph pairs by name."""
    from pptx import Presentation

    path, _ = exported
    deck = Presentation(str(path))

    names = [shape.name for slide in deck.slides for shape in slide.shapes]
    assert names, "the deck exported no shapes at all"
    assert all(name.startswith("deckastra-el_") for name in names)
    # A counter would renumber whenever a slide gains an element, and every
    # morph on the deck would silently stop pairing.
    assert len(set(names)) == len(names)


def test_speaker_notes_survive(exported):
    """A note written in the editor comes back out through a reader that is not ours.

    This used to assert only that the fixture's slides had *no* notes, which is
    true of a correct exporter and just as true of one writing them into an
    element no conforming reader looks for — which is what it was doing. The
    part is `notesSlides/notesSlide1.xml` and the element inside it is
    `<p:notes>`, not `<p:notesSlide>`; the two names differ, and nothing noticed
    while the only thing reading these files was the code that wrote them.

    `python-pptx` did not merely disagree: it raised, because it could not map
    the root element to its notes class at all.
    """
    from pptx import Presentation

    path, _ = exported
    deck = Presentation(str(path))

    with_notes = [slide for slide in deck.slides if slide.has_notes_slide]
    assert with_notes, "the deck exported no notes slide at all"

    text = with_notes[0].notes_slide.notes_text_frame.text
    assert "Remember to mention the budget." in text
    # Both lines: a note is paragraphs, and flattening them to one would lose
    # the shape of what somebody wrote.
    assert "And thank the team." in text

    # And only the slide that has one: a notes part per slide regardless would
    # bloat every deck and tell PowerPoint there is something to show.
    assert len(with_notes) == 1


def _all_shapes(shapes):
    for shape in shapes:
        yield shape
        if shape.shape_type == 6:  # MSO_SHAPE_TYPE.GROUP
            yield from _all_shapes(shape.shapes)


def test_a_table_arrives_as_a_real_table_with_its_cells(exported):
    """Tables used to be a labelled placeholder box: the numbers were not in the file."""
    from pptx import Presentation

    path, report = exported
    deck = Presentation(str(path))
    document = json.loads(FIXTURE.read_text(encoding="utf-8"))
    index = next(i for i, slide in enumerate(document["slides"]) if any(e["type"] == "table" for e in slide["elements"]))
    element = next(e for e in document["slides"][index]["elements"] if e["type"] == "table")

    tables = [shape.table for shape in deck.slides[index].shapes if shape.has_table]
    assert len(tables) == 1, "the table is not a PowerPoint table"
    table = tables[0]
    header = [table.cell(0, c).text for c in range(len(table.columns))]
    assert header == [column.get("label", "") for column in element["columns"]]

    def text(cell):
        content = cell["content"]
        return content if isinstance(content, str) else "\n".join("".join(s["text"] for s in b["spans"]) for b in content["blocks"])

    first_row = [table.cell(1, c).text for c in range(len(table.columns))]
    assert first_row[0] == text(element["rows"][0]["cells"][0])
    dropped = [w["feature"] for w in report["warnings"] if w["action"] == "dropped"]
    assert "table" not in dropped and "diagram" not in dropped, dropped


def test_a_diagram_arrives_as_shapes_with_its_labels_editable(exported):
    from pptx import Presentation

    path, _ = exported
    deck = Presentation(str(path))
    document = json.loads(FIXTURE.read_text(encoding="utf-8"))
    index = next(i for i, slide in enumerate(document["slides"]) if any(e["type"] == "diagram" for e in slide["elements"]))
    labels = [node["label"] for e in document["slides"][index]["elements"] if e["type"] == "diagram" for node in e["nodes"]]

    texts = [shape.text_frame.text for shape in _all_shapes(deck.slides[index].shapes) if shape.has_text_frame]
    for label in labels:
        assert any(label in text for text in texts), f"the diagram's {label!r} is not in the file"
