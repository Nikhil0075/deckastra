"""Build the eight-language visual benchmark as a real editable deck.

This is intentionally not a collage generator.  Each slide is composed by the
same design-language function used by templates, then supplied with a reviewed
reference image.  The resulting JSON, PDF and PowerPoint are therefore a useful
regression artifact: fixing it fixes the product rather than one screenshot.
"""

from __future__ import annotations

import base64
from copy import deepcopy
import hashlib
import html
import json
from pathlib import Path
import sys
from typing import Any

ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(ROOT / "apps" / "api"))

from PIL import Image  # noqa: E402

from deckastra_api import export_service, presets  # noqa: E402
from deckastra_api.compose import compose_document  # noqa: E402
from deckastra_api.ids import new_id  # noqa: E402
from deckastra_api.models import SlideLayout, SlidePlan, StoryPlan  # noqa: E402
from deckastra_api.schema import validate_document  # noqa: E402


MEDIA = ROOT / "packages" / "deck-presets" / "media"
OUT = ROOT / ".artifacts" / "design-system-reference-deck"

SYSTEMS: tuple[dict[str, Any], ...] = (
    {
        "language": "swiss-signal",
        "theme": "swiss-signal",
        "title": "SWISS SIGNAL",
        "eyebrow": "DESIGN SYSTEM · 01",
        "subtitle": "Structure creates the signal.",
        "media": (("Swiss architecture frame", "swiss-signal-architecture-v3.jpg", "Monochrome curved concrete architecture"),),
    },
    {
        "language": "cinema-noir",
        "theme": "cinema-noir",
        "title": "CINEMA NOIR",
        "eyebrow": "CASE FILE · 02",
        "subtitle": "A story told in shadow, evidence and consequence.",
        "media": (
            ("Image frame", "noir-night-story-scene.jpg", "A solitary figure in a rain-darkened city street"),
            ("Contact print", "noir-case-file-scene.jpg", "Noir case-file detail"),
            ("Contact print", "noir-event-story-scene.jpg", "Noir event detail"),
            ("Contact print", "noir-thesis-defence-scene.jpg", "Noir research detail"),
        ),
    },
    {
        "language": "play-lab",
        "theme": "playful-pastel",
        "title": "PLAY LAB",
        "eyebrow": "TEAM KICKOFF · 03",
        "subtitle": "Make the next move impossible to miss.",
        "media": (("Play Lab hero frame", "play-lab-motion-v2.jpg", "A runner leaping through vivid geometric forms"),),
    },
    {
        "language": "system-terminal",
        "theme": "system-terminal",
        "title": "SYSTEM TERMINAL",
        "eyebrow": "SYSTEM MAP · 04",
        "subtitle": "Inspect the layers. Follow the signal.",
        "media": (),
    },
    {
        "language": "quiet-luxe",
        "theme": "quiet-luxury",
        "title": "QUIET LUXE",
        "eyebrow": "COLLECTION · 05",
        "subtitle": "Material, restraint and a single considered gesture.",
        "media": (
            ("Portrait frame", "quiet-luxe-interior-v3.jpg", "Burgundy chair, marble and olive branches in a quiet interior"),
            ("Luxe inset", "luxe-business-pitch-hero.jpg", "A restrained material detail"),
        ),
    },
    {
        "language": "data-desk",
        "theme": "civic",
        "title": "DATA DESK",
        "eyebrow": "FIELD REPORT · 06",
        "subtitle": "One page for the evidence that changes the decision.",
        "media": (),
    },
    {
        "language": "earth-story",
        "theme": "forest",
        "title": "EARTH STORY",
        "eyebrow": "FIELD NOTES · 07",
        "subtitle": "A landscape of evidence, texture and lived detail.",
        "media": (
            ("Arch frame", "earth-story-mountain-v3.jpg", "Green mountains, cloud and waterfalls"),
            ("Earth detail frame", "earth-customer-case-study-scene.jpg", "A close natural field-study detail"),
            ("Earth texture frame", "earth-lesson-plan-hero.jpg", "A warm paper and botanical texture"),
            ("Earth texture frame", "earth-training-module-scene.jpg", "A field study in open country"),
        ),
    },
    {
        "language": "spatial-future",
        "theme": "glassmorphism",
        "title": "SPATIAL FUTURE",
        "eyebrow": "ORBITAL BRIEF · 08",
        "subtitle": "Turn complexity into a luminous field of view.",
        "media": (
            ("Hologram frame", "space-product-launch-scene.jpg", "Luminous orbital forms in deep space"),
            ("Spatial thumbnail", "space-research-talk-scene.jpg", "A spatial research scene"),
            ("Spatial thumbnail", "asterloop-closing-v2.jpg", "Earth at night from orbit"),
            ("Spatial thumbnail", "asterloop-roadmap-v2.jpg", "A glowing path through a mountain landscape"),
        ),
    },
)


def _token(theme: dict[str, Any], value: str) -> Any:
    cursor: Any = theme
    for key in value.removeprefix("token:").split("."):
        if not isinstance(cursor, dict) or key not in cursor:
            return value
        cursor = cursor[key]
    return deepcopy(cursor)


def _resolve(value: Any, theme: dict[str, Any]) -> Any:
    if isinstance(value, dict):
        return {key: _resolve(item, theme) for key, item in value.items()}
    if isinstance(value, list):
        return [_resolve(item, theme) for item in value]
    if isinstance(value, str) and value.startswith("token:"):
        resolved = _token(theme, value)
        return _resolve(resolved, theme) if resolved != value else value
    return value


def _compose(system: dict[str, Any]) -> tuple[dict[str, Any], dict[str, Any]]:
    theme, theme_id = presets.resolve_theme(system["theme"])
    plan = StoryPlan(
        title=system["title"],
        audience="Design review",
        objective="Match the approved visual-language reference",
        narrative_arc="One design language, rendered as a production slide.",
        slides=[SlidePlan(
            layout=SlideLayout.TITLE,
            purpose="Establish the design language",
            key_message=system["title"],
            headline=system["title"],
            eyebrow=system["eyebrow"],
            subtitle=system["subtitle"],
        )],
    )
    document = compose_document(
        plan,
        instruction="Eight design-system reference benchmark",
        theme_definition=theme,
        theme_id=theme_id,
        language=system["language"],
        language_version=presets.language_version(system["language"]),
        preserve_language_media=True,
    )
    return document, theme


def _asset(path: Path, alt: str) -> tuple[dict[str, Any], dict[str, Any]]:
    raw = path.read_bytes()
    with Image.open(path) as picture:
        width, height = picture.size
        has_alpha = picture.mode in {"RGBA", "LA"} and picture.getextrema()[-1][0] < 255
    identifier = new_id("ast")
    mime = "image/png" if path.suffix.lower() == ".png" else "image/jpeg"
    reference = {
        "id": identifier,
        "type": "image",
        "storageKey": f"benchmarks/design-system/{path.name}",
        "fileName": path.name,
        "mimeType": mime,
        "byteSize": len(raw),
        "width": width,
        "height": height,
        "checksum": hashlib.sha256(raw).hexdigest(),
        "altText": alt,
        "hasAlpha": has_alpha,
        "createdBy": "generated" if path.name.endswith("-v3.jpg") else "integration",
        "license": {"kind": "generated", "attribution": "Generated for Deckastra's release-safe benchmark"},
    }
    inline = {
        "assetId": identifier,
        "storageKey": reference["storageKey"],
        "mimeType": mime,
        "data": base64.b64encode(raw).decode("ascii"),
    }
    return reference, inline


def _place_media(
    slide: dict[str, Any],
    assignments: tuple[tuple[str, str, str], ...],
    assets: list[dict[str, Any]],
    inline: list[dict[str, Any]],
) -> None:
    occurrences: dict[str, int] = {}
    for frame_name, filename, alt in assignments:
        wanted = occurrences.get(frame_name, 0)
        occurrences[frame_name] = wanted + 1
        found = [
            (index, element)
            for index, element in enumerate(slide["elements"])
            if element.get("name") == frame_name
        ]
        if wanted >= len(found):
            raise RuntimeError(f"{slide['name']}: no occurrence {wanted + 1} of {frame_name!r}")
        index, frame = found[wanted]
        reference, payload = _asset(MEDIA / filename, alt)
        assets.append(reference)
        inline.append(payload)
        treatment: dict[str, Any] | None = None
        if frame_name in {"Image frame", "Contact print"}:
            treatment = {"monochrome": 1, "contrast": 0.18, "grain": 0.08, "vignette": 0.18}
        image: dict[str, Any] = {
            "id": new_id("el"),
            "type": "image",
            "name": f"{frame_name} image",
            "semanticRole": "supportingVisual",
            "layerRole": "foreground-subject" if frame_name == "Play Lab hero frame" else "content-media",
            "transform": deepcopy(frame["transform"]),
            "style": deepcopy(frame.get("style") or {}),
            "assetId": reference["id"],
            "fit": "cover",
            "altText": alt,
            "placeholder": "color",
            **({"treatment": treatment} if treatment else {}),
        }
        slide["elements"].insert(index + 1, image)


def _html(slides: list[Path]) -> str:
    cards = "\n".join(
        f'<figure><img src="slides/{html.escape(path.name)}" alt="{html.escape(SYSTEMS[index]["title"])}"><figcaption>{index + 1:02d} · {html.escape(SYSTEMS[index]["title"])}</figcaption></figure>'
        for index, path in enumerate(slides)
    )
    return f"""<!doctype html><meta charset=\"utf-8\"><title>Deckastra design-system benchmark</title>
<style>
*{{box-sizing:border-box}}body{{margin:0;background:#0d1117;color:#f5f7fa;font:15px Inter,Arial,sans-serif}}header{{padding:30px 34px 18px}}h1{{margin:0 0 8px;font-size:28px}}p{{margin:0;color:#aeb7c3}}main{{display:grid;grid-template-columns:repeat(2,minmax(0,1fr));gap:18px;padding:18px 34px 38px}}figure{{margin:0;background:#151b23;border:1px solid #303946;padding:10px}}img{{display:block;width:100%;aspect-ratio:16/9;object-fit:cover}}figcaption{{padding:10px 2px 2px;font-weight:700;letter-spacing:.08em}}@media(max-width:900px){{main{{grid-template-columns:1fr}}}}
</style><header><h1>Deckastra — eight design systems</h1><p>Actual renderer output · editable geometry · export benchmark</p></header><main>{cards}</main>"""


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    slide_dir = OUT / "slides"
    slide_dir.mkdir(parents=True, exist_ok=True)

    slides: list[dict[str, Any]] = []
    assets: list[dict[str, Any]] = []
    inline: list[dict[str, Any]] = []
    base_document: dict[str, Any] | None = None

    for index, system in enumerate(SYSTEMS, 1):
        composed, theme = _compose(system)
        if base_document is None:
            base_document = composed
        slide = _resolve(deepcopy(composed["slides"][0]), theme)
        slide["name"] = f"{index:02d} · {system['title']}"
        slide["background"] = _resolve(slide["background"], theme)
        _place_media(slide, system["media"], assets, inline)
        slides.append(slide)

    assert base_document is not None
    document = base_document
    document["metadata"] = {
        "title": "Deckastra Design Systems — Reference Benchmark",
        "description": "Eight production design languages matched against the approved reference collage.",
        "language": "en",
        "presentationType": "custom",
        "audience": "Deckastra design review",
        "objective": "Audit font, layout, colour and image placement across all eight systems.",
        "estimatedDurationSeconds": 480,
    }
    document["slides"] = slides
    document["assets"] = assets
    document["extensions"]["deckastra.designBenchmark"] = {
        "version": 1,
        "systems": [one["language"] for one in SYSTEMS],
        "reference": "approved eight-system collage",
    }

    errors = validate_document(document)
    if errors:
        raise RuntimeError("Benchmark document is invalid:\n" + "\n".join(errors))

    source = OUT / "design-system-reference.mydeck.json"
    source.write_text(json.dumps(document, ensure_ascii=False, indent=2), encoding="utf-8")

    pictures: list[Path] = []
    for index, slide in enumerate(slides, 1):
        result = export_service.render_slide_png(document, slide["id"], assets=inline)
        path = slide_dir / f"{index:02d}-{SYSTEMS[index - 1]['language']}.png"
        path.write_bytes(result["bytes"])
        pictures.append(path)
        if result["warnings"]:
            print(f"{path.name}: {result['warnings']}")

    for kind in ("pdf", "pptx"):
        target = OUT / f"design-system-reference.{kind}"
        export_service._invoke_worker(kind, document, target, {}, assets=inline)  # noqa: SLF001

    # Independent readers prove these are real files, not exporter success messages.
    from pypdf import PdfReader
    from pptx import Presentation

    pdf = PdfReader(str(OUT / "design-system-reference.pdf"))
    pptx = Presentation(str(OUT / "design-system-reference.pptx"))
    if len(pdf.pages) != 8 or len(pptx.slides) != 8:
        raise RuntimeError(f"Read-back failed: PDF={len(pdf.pages)} PPTX={len(pptx.slides)}")
    extracted = "\n".join(page.extract_text() or "" for page in pdf.pages)
    compact = "".join(extracted.split())
    missing = [
        system["title"]
        for system in SYSTEMS
        if "".join(system["title"].split()) not in compact
    ]
    if missing:
        raise RuntimeError("PDF text read-back missed: " + ", ".join(missing))

    (OUT / "index.html").write_text(_html(pictures), encoding="utf-8")
    (OUT / "audit.json").write_text(json.dumps({
        "slides": 8,
        "pdfPages": len(pdf.pages),
        "powerPointSlides": len(pptx.slides),
        "sourceImages": [one["fileName"] for one in assets],
        "checks": ["schema", "PNG render", "PDF read-back", "PowerPoint reopen", "PDF text extraction"],
    }, indent=2), encoding="utf-8")
    print(OUT)


if __name__ == "__main__":
    main()
