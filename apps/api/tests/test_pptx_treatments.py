"""A treated picture reopens in a PowerPoint reader that is not ours (2026-10-13).

`python-pptx` resolves the picture's relationship and exposes the blip it
embeds. Black and white must come back as PowerPoint's own greyscale effect,
and a duotone as its own duotone with the two colours, both on the picture
the reader found, so a person can still change them in Picture Format. A
reader that refused the file, or found no effect, would mean PowerPoint does
too.
"""

from __future__ import annotations

import json
import shutil
import subprocess
import sys
from pathlib import Path

import pytest

from tests.test_pptx_opens import ROOT, FIXTURE, _png

pptx = pytest.importorskip("pptx", reason="python-pptx is the independent reader this test needs")

if shutil.which("npx") is None:  # pragma: no cover - environment guard
    pytest.skip("npx is needed to run the exporter", allow_module_level=True)

EMIT = """
import { writeFileSync, readFileSync } from "node:fs";
import { buildDocumentScene } from "@deckastra/renderer";
import { fontManifest } from "@deckastra/export-core";
import { buildPptx } from "@deckastra/export-pptx";

const doc = JSON.parse(readFileSync(process.argv[2], "utf8"));
const image = doc.slides.flatMap((slide) => slide.elements).find((element) => element.type === "image");
image.treatment = JSON.parse(process.argv[5]);
const scene = buildDocumentScene(doc);
const scenes = new Map(scene.slides.map((s) => [s.slideId, s]));
const images = new Map([[image.assetId, { bytes: new Uint8Array(readFileSync(process.argv[4])), contentType: "image/png" }]]);
const { bytes, result } = buildPptx({ document: doc, scenes, fontManifest: fontManifest(scenes.values()), options: {}, images });
writeFileSync(process.argv[3], bytes);
process.stdout.write(JSON.stringify(result.report));
"""

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def _export(directory: Path, treatment: dict) -> Path:
    script = ROOT / "packages/export-pptx/.emit-treatments.mts"
    script.write_text(EMIT, encoding="utf-8")
    out = directory / "treated.pptx"
    picture = _png(directory / "picture.png")
    try:
        finished = subprocess.run(
            ["npx", "tsx", str(script), str(FIXTURE), str(out), str(picture), json.dumps(treatment)],
            cwd=ROOT,
            capture_output=True,
            text=True,
            timeout=180,
            shell=sys.platform == "win32",
        )
    finally:
        script.unlink(missing_ok=True)
    if finished.returncode != 0:
        pytest.fail(f"the exporter failed:\n{finished.stderr[-2000:]}")
    return out


def _blip(path: Path):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    deck = Presentation(str(path))
    pictures = [shape for slide in deck.slides for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
    assert pictures, "the treated picture is not in the file"
    picture = pictures[0]
    # The reader resolves the picture's bytes itself, so the relationship holds.
    assert picture.image.blob
    return picture._element.find(f".//{A}blip")


@pytest.mark.slow
def test_black_and_white_is_powerpoints_own_greyscale(tmp_path):
    blip = _blip(_export(tmp_path, {"monochrome": 1}))
    assert blip.find(f"{A}grayscl") is not None


@pytest.mark.slow
def test_a_duotone_is_powerpoints_own_duotone_with_both_colours(tmp_path):
    blip = _blip(_export(tmp_path, {"duotone": ["#102030", "#F0E0D0"]}))
    duotone = blip.find(f"{A}duotone")
    assert duotone is not None
    assert [colour.get("val") for colour in duotone.findall(f"{A}srgbClr")] == ["102030", "F0E0D0"]
